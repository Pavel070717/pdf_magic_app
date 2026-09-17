"""
Tests for routes/magic.py — file copy/numbering/registry endpoints.
"""

import json
from unittest.mock import MagicMock, patch

from routes.magic import (
    can_share_number,
    copy_files_worker,
    generate_numbered_filename,
    strip_leading_number,
)


class TestGenerateNumberedFilename:
    def test_basic(self):
        result = generate_numbered_filename(1, "document", ".pdf")
        assert result == "01.document.pdf"

    def test_double_digit(self):
        result = generate_numbered_filename(42, "report", ".docx")
        assert result == "42.report.docx"

    def test_triple_digit(self):
        result = generate_numbered_filename(100, "file", ".txt")
        assert result == "100.file.txt"

    def test_zero_padded(self):
        result = generate_numbered_filename(5, "name", ".ext")
        assert result == "05.name.ext"


class TestStripLeadingNumber:
    def test_single_prefix(self):
        assert strip_leading_number("01.Акт.pdf") == "Акт.pdf"

    def test_double_prefix(self):
        assert strip_leading_number("01.01.Акт.pdf") == "Акт.pdf"

    def test_multi_part_prefix(self):
        assert strip_leading_number("12.3.Акт") == "Акт"

    def test_year_not_stripped(self):
        assert strip_leading_number("2024.Отчет.pdf") == "2024.Отчет.pdf"

    def test_no_prefix(self):
        assert strip_leading_number("Акт") == "Акт"

    def test_prefix_with_space(self):
        assert strip_leading_number("02. Акт") == "Акт"


class TestCanShareNumber:
    def test_pdf_xlsx(self):
        assert can_share_number(".pdf", ".xlsx") is True

    def test_xlsx_pdf(self):
        assert can_share_number(".xlsx", ".pdf") is True

    def test_pdf_xlsm(self):
        assert can_share_number(".pdf", ".xlsm") is True

    def test_case_insensitive(self):
        assert can_share_number(".PDF", ".XLSX") is True

    def test_not_shared(self):
        assert can_share_number(".pdf", ".docx") is False

    def test_pdf_pdf_not_shared(self):
        assert can_share_number(".pdf", ".pdf") is False


class TestStartMagic:
    def test_no_files_in_state(self, app_client, clean_state):
        resp = app_client.post(
            "/api/magic/start",
            data=json.dumps({}),
            content_type="application/json",
        )
        assert resp.status_code == 400
        assert "Нет файлов" in resp.get_json()["error"]

    def test_nonexistent_target_dir(self, app_client, clean_state):
        resp = app_client.post(
            "/api/magic/start",
            data=json.dumps({"target_dir": "/nonexistent/path"}),
            content_type="application/json",
        )
        assert resp.status_code == 400

    def test_existing_numbered_files(self, app_client, clean_state, temp_dir):
        from utils.state import load_state, save_state

        state = load_state()
        state["files"] = [{"path": "/fake/file.pdf", "name": "file.pdf"}]
        save_state(state)
        (temp_dir / "01.existing.pdf").touch()
        resp = app_client.post(
            "/api/magic/start",
            data=json.dumps({"target_dir": str(temp_dir)}),
            content_type="application/json",
        )
        assert resp.status_code == 400
        assert "пронумерованные" in resp.get_json()["error"]

    @patch("routes.magic.threading.Thread")
    def test_success_starts_thread(
        self, mock_thread, app_client, clean_state, temp_dir
    ):
        from utils.state import load_state, save_state

        state = load_state()
        state["files"] = [{"path": "/fake/file.pdf", "name": "file.pdf"}]
        save_state(state)
        mock_thread.return_value = MagicMock()
        resp = app_client.post(
            "/api/magic/start",
            data=json.dumps({"target_dir": str(temp_dir)}),
            content_type="application/json",
        )
        assert resp.status_code == 200
        assert resp.get_json()["success"] is True

    def test_invalid_target_dir_fallback(self, app_client, clean_state):
        from utils.state import load_state, save_state

        state = load_state()
        state["files"] = [{"path": "/fake/file.pdf", "name": "file.pdf"}]
        save_state(state)
        resp = app_client.post(
            "/api/magic/start",
            data=json.dumps({"target_dir": "/nonexistent/dir"}),
            content_type="application/json",
        )
        assert resp.status_code in (200, 400)


class TestCopyWorker:
    def test_comma_in_name_survives_copy_and_registry(
        self, app_client, clean_state, temp_dir
    ):
        from openpyxl import load_workbook

        from utils.state import load_state, save_state

        src_dir = temp_dir / "src"
        src_dir.mkdir()
        src = src_dir / "15,1 м.pdf"
        src.write_bytes(b"%PDF-1.4\nplaceholder\n%%EOF")

        state = load_state()
        state["files"] = [
            {
                "id": "f1",
                "path": str(src),
                "name": src.name,
                "original_path": str(src),
            }
        ]
        save_state(state)

        target = temp_dir / "out"
        target.mkdir()
        copy_files_worker(state["files"], target)

        numbered = target / "01.15,1 м.pdf"
        assert numbered.exists(), sorted(p.name for p in target.iterdir())

        registry_files = list(target.glob("00.Реестр*.xlsx"))
        assert registry_files, sorted(p.name for p in target.iterdir())
        wb = load_workbook(registry_files[0], read_only=True)
        ws = wb[wb.sheetnames[0]]
        assert ws["B29"].value == "15,1 м"
        wb.close()

    def test_aosr_xlsx_not_copied_but_name_used(
        self, app_client, clean_state, temp_dir
    ):
        from openpyxl import Workbook, load_workbook

        from utils.state import load_state, save_state

        src_dir = temp_dir / "src"
        src_dir.mkdir()

        xlsx_path = src_dir / "АОСР_№1_22.05.2025.xlsx"
        book = Workbook()
        book.active["A77"] = "Бетонирование плиты"
        book.save(xlsx_path)
        book.close()

        pdf_path = src_dir / "АОСР_№1_22.05.2025.pdf"
        pdf_path.write_bytes(b"%PDF-1.4\nplaceholder\n%%EOF")
        act_path = src_dir / "Акт.pdf"
        act_path.write_bytes(b"%PDF-1.4\nplaceholder\n%%EOF")

        state = load_state()
        state["files"] = [
            {
                "id": "f1",
                "path": str(pdf_path),
                "name": pdf_path.name,
                "original_path": str(pdf_path),
            },
            {
                "id": "f2",
                "path": str(xlsx_path),
                "name": xlsx_path.name,
                "original_path": str(xlsx_path),
            },
            {
                "id": "f3",
                "path": str(act_path),
                "name": act_path.name,
                "original_path": str(act_path),
            },
        ]
        save_state(state)

        target = temp_dir / "out"
        target.mkdir()
        copy_files_worker(state["files"], target)

        names = {p.name for p in target.iterdir()}
        assert "01.АОСР_№1_22.05.2025.pdf" in names, names
        assert "02.Акт.pdf" in names, names
        assert not any(n.startswith("01.АОСР") and n.endswith(".xlsx") for n in names)

        registry_files = list(target.glob("00.Реестр*.xlsx"))
        assert registry_files, names
        wb = load_workbook(registry_files[0], read_only=True)
        ws = wb[wb.sheetnames[0]]
        aosr_b_cells = [
            ws[f"B{r}"].value
            for r in range(29, 35)
            if ws[f"B{r}"].value and str(ws[f"B{r}"].value).upper().startswith("АОСР")
        ]
        assert aosr_b_cells == ["АОСР. Бетонирование плиты"]
        wb.close()


class TestCancelMagic:
    def test_cancel(self, app_client):
        resp = app_client.post(
            "/api/magic/cancel",
            content_type="application/json",
        )
        assert resp.status_code == 200
        assert resp.get_json()["success"] is True


class TestMagicProgress:
    def test_get_progress(self, app_client):
        resp = app_client.get("/api/magic/progress")
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["success"] is True
        assert "progress" in data


class TestMagicResult:
    def test_get_result(self, app_client, clean_state):
        resp = app_client.get("/api/magic/result")
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["success"] is True
        assert "last_magic_run" in data
        assert "files" in data
